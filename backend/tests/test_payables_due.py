"""🆕 2026-09-29 采购「应付到期」表（GET /api/purchase-mgmt/payables-due）。

起因：有人提「账期比如 30 天，货到了以后要提醒该付款了，给采购部一张表」。
老板拍板：从**到货**那天算；先不推送，只要这张表。

本文件钉死：
  · 到期日 = 到货日期 + 供应商账期天数；分档（已过期/7天内/30天内/更晚/账期未填）正确；
  · 已付清的、没到货的不进表；账期未填的单独一档，不猜成「已过期」；
  · **和财务资金面板的「已过期」合计一分不差**（两边共用 app/payables.py）；
  · 受限采购员只看自己下的单；
  · 请款走到哪一步（未请款 / 请款审批中 / 已批待付款）标得对。
"""
import asyncio, os, sys, tempfile
from datetime import datetime, timedelta

tmp = tempfile.mkdtemp(prefix="apdue")
os.environ["DATABASE_URL"] = f"sqlite+aiosqlite:///{tmp}/test.db"
os.environ["FILES_DIR"] = f"{tmp}/files"
os.chdir(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.getcwd())

from httpx import AsyncClient, ASGITransport
from app.main import app
from app.database import engine, SessionLocal, Base
from app.seed import seed
from app.data_migration import run_all, ensure_schema_columns
from app import models
from app.overdue import _CN_TZ

FAIL = []


def chk(c, m):
    print(("  PASS " if c else "  FAIL: ") + m)
    if not c:
        FAIL.append(m)


TODAY = datetime.now(_CN_TZ).date()


def d(n):  # 相对北京时间今天的日期串
    return (TODAY + timedelta(days=n)).isoformat()


async def main():
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    await ensure_schema_columns(engine)
    async with SessionLocal() as db:
        await seed(db)
        await run_all(db)

    transport = ASGITransport(app=app, raise_app_exceptions=False)
    async with AsyncClient(transport=transport, base_url="http://test", timeout=60) as c:
        async def login(u, p="pass123"):
            r = await c.post("/api/auth/login", json={"username": u, "password": p})
            assert r.status_code == 200, r.text
            return {"Authorization": f"Bearer {r.json()['access_token']}"}

        H = await login("admin", "admin123")
        rid = {x["code"]: x["id"] for x in (await c.get("/api/admin/roles", headers=H)).json()}

        async def mkuser(name, codes):
            r = await c.post("/api/admin/users", headers=H, json={
                "username": name, "password": "pass123", "full_name": name,
                "role_ids": [rid[x] for x in codes]})
            assert r.status_code == 200, r.text
            return r.json()["id"]

        b1 = await mkuser("ap_buyer1", ["buyer"])
        b2 = await mkuser("ap_buyer2", ["buyer"])
        await mkuser("ap_fin", ["finance"])
        await mkuser("ap_asm", ["assembler"])
        Hb1, Hfin, Hasm = await login("ap_buyer1"), await login("ap_fin"), await login("ap_asm")

        # ── 造数据 ──
        async with SessionLocal() as db:
            s30 = models.Supplier(name="月结三十供应商", settlement_type="月结", credit_days=30)
            s60 = models.Supplier(name="月结六十供应商", settlement_type="月结", credit_days=60)
            scash = models.Supplier(name="现金供应商", settlement_type="现金", credit_days=None)
            db.add_all([s30, s60, scash])
            await db.flush()

            def item(sup, name, recv, paid, arrival, buyer, po):
                it = models.PurchaseItem(supplier_id=sup.id, item_name=name, qty=1,
                                         unit_price=recv, received_amount=recv, paid_amount=paid,
                                         arrival_date=arrival, buyer_id=buyer, po_no=po,
                                         project_code="2026-090")
                db.add(it)
                return it

            i_over = item(s30, "已过期的料", 1000, 0, d(-40), b1, "PO-A")        # 到期 d(-10)
            i_week = item(s30, "七天内的料", 2000, 500, d(-25), b1, "PO-B")      # 到期 d(+5)，欠 1500
            i_month = item(s60, "三十天内的料", 3000, 0, d(-40), b2, "PO-C")      # 到期 d(+20)
            i_later = item(s60, "更晚的料", 4000, 0, d(-5), b2, "PO-D")          # 到期 d(+55)
            i_cash = item(scash, "现金供应商的料", 500, 0, d(-90), b1, "PO-E")    # 账期未填
            item(s30, "已付清的料", 800, 800, d(-100), b1, "PO-F")               # 不该出现
            # 期初余额里已经含了的：腾丰那种（下单日期 ≤ 期初日期）。两边都不该算
            s_ob = models.Supplier(name="有期初的供应商", settlement_type="月结", credit_days=30)
            db.add(s_ob)
            await db.flush()
            db.add(models.SupplierOpeningBalance(supplier_id=s_ob.id, balance_date=d(-60),
                                                 outstanding_amount=12492))
            ob_item = models.PurchaseItem(supplier_id=s_ob.id, item_name="期初已含的料", qty=1,
                                          unit_price=12492, received_amount=12492, paid_amount=0,
                                          delivery_date=d(-90), arrival_date=d(-85),
                                          buyer_id=b1, po_no="PO-OB")
            db.add(ob_item)                                                       # 不该出现
            not_arr = models.PurchaseItem(supplier_id=s30.id, item_name="没到货的料", qty=1,
                                          unit_price=900, received_amount=0, paid_amount=0,
                                          arrival_date=None, buyer_id=b1, po_no="PO-G")
            db.add(not_arr)                                                       # 不该出现
            await db.flush()

            # 「七天内的料」已经提了请款，还在审批
            pr = models.PaymentRequest(supplier_id=s30.id, requested_amount=1500,
                                       status="pending", requester_id=b1)
            db.add(pr)
            await db.flush()
            db.add(models.PaymentRequestItem(request_id=pr.id, item_id=i_week.id,
                                             allocated_amount=1500))
            await db.commit()
            ids = {"over": i_over.id, "week": i_week.id, "month": i_month.id,
                   "later": i_later.id, "cash": i_cash.id}

        print("\n=== 1. 分档与到期日 ===")
        r = await c.get("/api/purchase-mgmt/payables-due", headers=H)
        chk(r.status_code == 200, f"管理层取表 → {r.status_code} {r.text[:80]}")
        body = r.json()
        rows = {x["item_id"]: x for x in body["rows"]}
        chk(set(rows) == set(ids.values()),
            f"只有「已到货、没付清」的进表（已付清的、没到货的、**期初余额里已含的**都不在）→ {sorted(rows)}")
        chk(not any(x["item_name"] == "期初已含的料" for x in body["rows"]),
            "期初已含的不重复算（与供应商账目同口径；腾丰那 ¥12,492 就是这么多出来的）")
        chk(rows[ids["over"]]["bucket"] == "overdue" and rows[ids["over"]]["due_date"] == d(-10)
            and rows[ids["over"]]["days_left"] == -10,
            f"到货 40 天前 + 账期 30 天 = 10 天前到期，已过期 → {rows[ids['over']]}")
        chk(rows[ids["week"]]["bucket"] == "week" and rows[ids["week"]]["outstanding"] == 1500,
            "7 天内到期；欠款 = 收货 2000 − 已付 500 = 1500")
        chk(rows[ids["month"]]["bucket"] == "month", "30 天内到期")
        chk(rows[ids["later"]]["bucket"] == "later", "30 天以后")
        chk(rows[ids["cash"]]["bucket"] == "nocredit" and rows[ids["cash"]]["due_date"] is None,
            "供应商没填账期天数 → 「账期未填」，**不猜成已过期**")
        order = [x["item_id"] for x in body["rows"]]
        chk(order == [ids["over"], ids["week"], ids["month"], ids["later"], ids["cash"]],
            f"按到期日排：最早该付的在最上面，账期未填的在最后 → {order}")
        sm = body["summary"]
        chk(sm["overdue"]["amount"] == 1000 and sm["week"]["amount"] == 1500
            and sm["nocredit"]["count"] == 1,
            f"汇总数对 → {sm}")

        print("\n=== 2. 请款状态 ===")
        chk(rows[ids["week"]]["request_state"] == "请款审批中", "提了请款还在审批的标出来")
        chk(rows[ids["over"]]["request_state"] == "未请款", "没提请款的标「未请款」")

        print("\n=== 3. 和财务资金面板对得上 ===")
        fp = await c.get("/api/reports/fund-panel", headers=Hfin)
        chk(fp.status_code == 200, f"资金面板 → {fp.status_code}")
        fp_over = round(fp.json()["payables"]["overdue_total"], 2)
        rf = (await c.get("/api/purchase-mgmt/payables-due", headers=Hfin)).json()
        chk(rf["summary"]["overdue"]["amount"] == fp_over,
            f"**已过期合计 = 资金面板的已过期合计**（两边共用 app/payables.py）"
            f"→ 表 {rf['summary']['overdue']['amount']} vs 面板 {fp_over}")
        chk(fp_over == 1000.0,
            f"资金面板也不再把期初已含的 12492 算成已过期（改之前是 13492）→ {fp_over}")
        fp_due_soon = round(fp.json()["payables"]["due_soon_total"], 2)
        tbl_14 = round(sum(x["outstanding"] for x in rf["rows"]
                           if x["days_left"] is not None and 0 <= x["days_left"] <= 14), 2)
        chk(fp_due_soon == tbl_14,
            f"14 天内到期也对得上 → 面板 {fp_due_soon} vs 表 {tbl_14}")

        print("\n=== 3b. 导出 Excel：屏幕上看到什么，导出来就是什么 ===")
        import io
        from openpyxl import load_workbook
        rx = await c.get("/api/purchase-mgmt/payables-due/export", headers=Hfin)
        chk(rx.status_code == 200 and "spreadsheetml" in rx.headers.get("content-type", ""),
            f"导出全部 → {rx.status_code} {rx.headers.get('content-type')}")
        ws = load_workbook(io.BytesIO(rx.content)).active
        data_rows = [r for r in ws.iter_rows(min_row=3, values_only=True) if r[3] and r[0] != "合计"]
        chk(len(data_rows) == 5, f"5 条明细都在 → {len(data_rows)}")
        chk(ws.cell(row=2, column=15).value == "欠款", "表头有「欠款」列")
        rx2 = await c.get("/api/purchase-mgmt/payables-due/export",
                          params={"bucket": "overdue"}, headers=Hfin)
        ws2 = load_workbook(io.BytesIO(rx2.content)).active
        rows2 = [r for r in ws2.iter_rows(min_row=3, values_only=True) if r[3] and r[0] != "合计"]
        chk(len(rows2) == 1 and rows2[0][8] == "已过期的料",
            f"只导「已过期」时就只有那一条 → {[r[8] for r in rows2]}")
        rx3 = await c.get("/api/purchase-mgmt/payables-due/export",
                          params={"q": "月结六十"}, headers=Hfin)
        ws3 = load_workbook(io.BytesIO(rx3.content)).active
        rows3 = [r for r in ws3.iter_rows(min_row=3, values_only=True) if r[3] and r[0] != "合计"]
        chk(len(rows3) == 2, f"按供应商关键字筛 → {len(rows3)} 条")
        rx4 = await c.get("/api/purchase-mgmt/payables-due/export", headers=Hb1)
        ws4 = load_workbook(io.BytesIO(rx4.content)).active
        rows4 = [r for r in ws4.iter_rows(min_row=3, values_only=True) if r[3] and r[0] != "合计"]
        chk(len(rows4) == 3, f"受限采购员导出也只有自己的 3 条（导出不能绕过行级隔离）→ {len(rows4)}")

        print("\n=== 4. 权限 ===")
        r1 = (await c.get("/api/purchase-mgmt/payables-due", headers=Hb1)).json()
        mine = {x["item_id"] for x in r1["rows"]}
        chk(mine == {ids["over"], ids["week"], ids["cash"]},
            f"受限采购员只看自己下的单（与采购明细同口径）→ {sorted(mine)}")
        chk(len(rf["rows"]) == 5, "财务看全部")
        ra = await c.get("/api/purchase-mgmt/payables-due", headers=Hasm)
        chk(ra.status_code == 403, f"装配工无权 → {ra.status_code}")


asyncio.run(main())
print("\n" + ("全部通过 ✅" if not FAIL else f"{len(FAIL)} 条失败 ❌"))
for m in FAIL:
    print(" -", m)
sys.exit(1 if FAIL else 0)
