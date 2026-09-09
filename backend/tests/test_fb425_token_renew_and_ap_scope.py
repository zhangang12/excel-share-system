"""反馈#425（卢照坤）登录自动续期 + 应付口径两项决策（2026-09-09 老板定）。

#425「每天下午四点左右系统自动退出登录」：令牌 8 小时到期、没有续期机制。口径：
  · 剩余有效期 < 一半 → 响应头 X-PMS-Refresh-Token 下发新令牌（deps.get_current_user）
  · 剩余充足 → 不发（别每个请求都签发）
  · 新令牌能用；旧令牌到期前也仍能用
决策 1「到货才算应付」：账目一览/对账单导出/汇总 KPI/明细页欠款 只算已到货的明细；未到货的订单额与预付款单列。
决策 2「期初余额只累计 balance_date 之后的明细」：期初之前下单的明细，收货/已付都不再累计（否则与期初重复）。
"""
import asyncio, os, sys, tempfile

tmp = tempfile.mkdtemp(prefix="fb425")
os.environ["DATABASE_URL"] = os.environ.get("AUDIT_TEST_DATABASE_URL") or f"sqlite+aiosqlite:///{tmp}/test.db"
os.environ["FILES_DIR"] = f"{tmp}/files"
os.chdir(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.getcwd())

from httpx import AsyncClient, ASGITransport
from sqlalchemy import select
from app.main import app
from app.database import engine, SessionLocal, Base
from app.seed import seed
from app.data_migration import run_all, ensure_schema_columns
from app import models
from app.auth import create_access_token, decode_token

FAIL = []


def chk(c, m):
    print(("  PASS " if c else "  FAIL: ") + m)
    if not c:
        FAIL.append(m)


async def main():
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    await ensure_schema_columns(engine)
    async with SessionLocal() as db:
        await seed(db)
        await run_all(db)

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test", timeout=60) as c:
        async def login(u, p="pass123"):
            r = await c.post("/api/auth/login", json={"username": u, "password": p})
            assert r.status_code == 200, r.text
            return {"Authorization": f"Bearer {r.json()['access_token']}"}

        H = await login("admin", "admin123")
        rid = {x["code"]: x["id"] for x in (await c.get("/api/admin/roles", headers=H)).json()}

        async def mkuser(name, codes):
            r = await c.post("/api/admin/users", headers=H, json={
                "username": name, "password": "pass123", "full_name": name,
                "role_ids": [rid[x] for x in codes]})
            assert r.status_code == 200, r.text
            return r.json()["id"]

        wh_id = await mkuser("fb425_wh", ["warehouse"])
        await mkuser("fb425_buyer", ["buyer"])
        Hb = await login("fb425_buyer")
        Hw = await login("fb425_wh")

        # ══════════════ ① 登录续期 ══════════════
        print("\n① 登录自动续期（#425）")
        r = await c.get("/api/wh/materials", headers=Hw)
        chk(r.status_code == 200 and "x-pms-refresh-token" not in r.headers,
            f"刚登录（剩余 8 小时）→ 不下发新令牌 -> {r.status_code} {'x-pms-refresh-token' in r.headers}")
        # 造一张只剩 30 分钟的令牌（相当于下午 15:30 的仓库）
        old = create_access_token(wh_id, minutes=30)
        r = await c.get("/api/wh/materials", headers={"Authorization": f"Bearer {old}"})
        new = r.headers.get("x-pms-refresh-token")
        chk(r.status_code == 200 and bool(new), f"剩余 30 分钟 → 响应头带新令牌 -> {r.status_code} {bool(new)}")
        p_old, p_new = decode_token(old), decode_token(new or "")
        chk(bool(p_new) and str(p_new.get("sub")) == str(wh_id), "新令牌是同一个人")
        chk(bool(p_new) and p_new["exp"] - p_old["exp"] > 7 * 3600, f"新令牌有效期比旧的长 ≥7 小时 -> {(p_new or {}).get('exp', 0) - p_old['exp']}s")
        r2 = await c.get("/api/wh/materials", headers={"Authorization": f"Bearer {new}"})
        chk(r2.status_code == 200 and "x-pms-refresh-token" not in r2.headers, "新令牌能用，且不会再立刻续一张")
        r3 = await c.get("/api/wh/materials", headers={"Authorization": f"Bearer {old}"})
        chk(r3.status_code == 200, "旧令牌到期前仍能用（前端换令牌失败也不会立刻掉线）")
        r4 = await c.get("/api/wh/materials", headers={"Authorization": f"Bearer {create_access_token(wh_id, minutes=-1)}"})
        chk(r4.status_code == 401, f"已过期的令牌仍是 401（续期不是无限延寿）-> {r4.status_code}")

        # ══════════════ ② 应付口径 ══════════════
        print("\n② 到货才算应付 + 期初日期")
        sid = (await c.post("/api/purchase-mgmt/suppliers", headers=Hb, json={"name": "腾丰式供应商"})).json()["id"]
        r = await c.post(f"/api/purchase-mgmt/suppliers/{sid}/opening-balance", headers=Hb,
                         json={"balance_date": "2026-08-06", "outstanding_amount": 1000})
        chk(r.status_code == 200, f"设期初 2026-08-06 欠 1000 -> {r.status_code} {r.text[:80]}")

        async def order(name, price, delivery_date, arrive=False, pay=0):
            r = await c.post("/api/purchase-mgmt/orders", headers=Hb, json={
                "supplier_id": sid, "delivery_date": delivery_date,
                "lines": [{"item_name": name, "qty": 1, "unit_price": price}]})
            assert r.status_code == 200, r.text
            iid = r.json()[0]["id"]
            if arrive:
                rr = await c.put(f"/api/purchase-mgmt/items/{iid}/receive", headers=Hw, json={"arrival_date": "2026-09-01"})
                assert rr.status_code == 200, rr.text
            if pay:
                rr = await c.put(f"/api/purchase-mgmt/items/{iid}", headers=Hb, json={"paid_amount": pay, "paid_date": "2026-09-02"})
                assert rr.status_code == 200, rr.text
            return iid

        await order("期初前·已到货", 300, "2026-07-29", arrive=True)          # 含在期初里 → 不累计
        await order("期初后·已到货", 500, "2026-08-20", arrive=True, pay=200)  # 应付：500 收 / 200 付
        await order("期初后·未到货", 800, "2026-09-05")                        # 未到货订单 800，不算应付
        await order("期初后·未到货已预付", 400, "2026-09-06", pay=400)          # 预付 400，不算应付、不抵欠款

        st = (await c.get("/api/purchase-mgmt/statements", headers=Hb)).json()
        row = next(x for x in st["rows"] if x["supplier_id"] == sid)
        chk(row["received_total"] == 500, f"收货合计只算期初后·已到货 = 500（期初前 300 不累计）-> {row['received_total']}")
        chk(row["paid_total"] == 200, f"已付只算已到货那条 = 200 -> {row['paid_total']}")
        chk(row["outstanding"] == 1300, f"欠款 = 期初 1000 + 500 − 200 = 1300（原口径会是 1000+2000−600=2400）-> {row['outstanding']}")
        chk(row["pending_total"] == 1200 and row["prepaid_total"] == 400,
            f"未到货订单 1200 / 预付 400 单列 -> {row['pending_total']}/{row['prepaid_total']}")
        chk(row["item_count"] == 4, "明细数仍是全部 4 条")
        chk(st["total_prepaid"] == 400 and "到货" in (st.get("note") or ""), "合计与口径说明都带上")

        kpi = (await c.get("/api/purchase-mgmt/reports/overview", headers=Hb)).json()
        chk(abs(kpi["total_outstanding"] - 1300) < 0.01, f"汇总 KPI 欠款同口径 = 1300 -> {kpi['total_outstanding']}")

        summ = (await c.get("/api/purchase-mgmt/items/summary", headers=Hb, params={"supplier_id": sid})).json()
        chk(summ["received_total"] == 2000 and summ["outstanding"] == 600,
            f"明细页：收货合计仍是列表总额 2000；欠款只算已到货 800−200 = 600 -> {summ['received_total']}/{summ['outstanding']}")

        r = await c.get(f"/api/purchase-mgmt/statements/{sid}/export", headers=Hb)
        chk(r.status_code == 200 and len(r.content) > 1000, f"对账单导出正常 -> {r.status_code}")
        from openpyxl import load_workbook
        from io import BytesIO
        ws = load_workbook(BytesIO(r.content), read_only=True).active
        summ_line = str(ws.cell(row=3, column=1).value)
        chk("欠款余额：1,300.00" in summ_line and "预付(未到货)：400.00" in summ_line,
            f"导出汇总条同口径 -> {summ_line[:120]}")

        # 没有期初余额的供应商：只看到没到货
        sid2 = (await c.post("/api/purchase-mgmt/suppliers", headers=Hb, json={"name": "无期初供应商"})).json()["id"]
        r = await c.post("/api/purchase-mgmt/orders", headers=Hb, json={
            "supplier_id": sid2, "delivery_date": "2026-01-01",
            "lines": [{"item_name": "老单子已到货", "qty": 2, "unit_price": 50}]})
        iid = r.json()[0]["id"]
        await c.put(f"/api/purchase-mgmt/items/{iid}/receive", headers=Hw, json={"arrival_date": "2026-01-05"})
        st = (await c.get("/api/purchase-mgmt/statements", headers=Hb)).json()
        row2 = next(x for x in st["rows"] if x["supplier_id"] == sid2)
        chk(row2["received_total"] == 100 and row2["outstanding"] == 100, f"无期初的供应商不受日期过滤 -> {row2['outstanding']}")

    print()
    if FAIL:
        print(f"❌ {len(FAIL)} 项失败")
        sys.exit(1)
    print("✅ 全部通过")


asyncio.run(main())
